import json
import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parent
OUT_DOCX = BASE_DIR / "Informe_Proyecto_Rutas_Recoleccion.docx"


def extract_const_from_html(path: Path, name: str):
    txt = path.read_text(encoding="utf-8", errors="ignore")
    match = re.search(rf"const\s+{name}\s*=\s*", txt)
    if not match:
        raise ValueError(f"No se encontro const {name} en {path.name}")

    start = match.end()
    depth = 0
    in_string = False
    escape = False

    for pos in range(start, len(txt)):
        ch = txt[pos]
        if in_string:
            if escape:
                escape = False
            elif ch == "\\":
                escape = True
            elif ch == '"':
                in_string = False
        else:
            if ch == '"':
                in_string = True
            elif ch in "[{":
                depth += 1
            elif ch in "]}":
                depth -= 1
            elif ch == ";" and depth == 0:
                return json.loads(txt[start:pos])

    raise ValueError(f"No se pudo cerrar const {name} en {path.name}")


def fmt_num(value, digits=2):
    if value is None:
        return "-"
    try:
        return f"{float(value):,.{digits}f}"
    except (TypeError, ValueError):
        return str(value)


def fmt_money(value):
    return f"${fmt_num(value, 2)}"


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(str(text))
    run.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        set_cell_text(hdr[i], header, bold=True)
        shading = hdr[i]._tc.get_or_add_tcPr()
        # Keep default Word styling; avoid brittle custom XML here.
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    doc.add_paragraph()
    return table


def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_paragraph(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def metricas_escenario(nombre, html_name):
    html = BASE_DIR / html_name
    metricas = extract_const_from_html(html, "METRICAS")
    clientes = extract_const_from_html(html, "CLIENTES")
    return {
        "nombre": nombre,
        "archivo": html_name,
        "global": metricas["global"],
        "camiones": metricas["camiones"],
        "viajes": metricas["viajes"],
        "clientes_features": len(clientes["features"]),
    }


def km_viaje(v):
    return sum(float(v.get(k, 0.0)) for k in ("aprox_km", "recoleccion_km", "descarga_km", "cierre_km"))


def min_viaje(v):
    return sum(
        float(v.get(k, 0.0))
        for k in (
            "aprox_min",
            "recoleccion_min",
            "espera_desalojo_min",
            "descarga_min",
            "cierre_min",
            "balance_turno_min",
        )
    )


def crear_informe():
    esc60 = metricas_escenario("Escenario 60 toneladas", "mapa_rutas_recoleccion_60t.html")
    esc120 = metricas_escenario("Escenario 120 toneladas", "mapa_rutas_recoleccion_120t.html")

    comparacion_alg = pd.read_csv(BASE_DIR / "comparacion_algoritmos.csv")
    comparacion_k = pd.read_csv(BASE_DIR / "comparacion_k_zonas.csv")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        styles[style_name].font.name = "Calibri"
        styles[style_name].font.color.rgb = RGBColor(31, 78, 121)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Informe del Proyecto de Optimización de Rutas de Recolección de Residuos")
    run.bold = True
    run.font.size = Pt(18)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Escenarios operativos de 60 toneladas y 120 toneladas")
    r.italic = True
    r.font.size = Pt(13)

    doc.add_paragraph()
    add_table(
        doc,
        ["Campo", "Detalle"],
        [
            ["Proyecto", "Optimización y visualización de rutas de recolección urbana"],
            ["Área", "Redes / análisis geoespacial / logística urbana"],
            ["Lugar de estudio", "Cuenca, Ecuador"],
            ["Autor(es)", "Completar con nombres del equipo"],
            ["Fecha", "Junio de 2026"],
        ],
    )

    add_heading(doc, "1. Resumen ejecutivo")
    add_paragraph(
        doc,
        "El proyecto implementa un flujo computacional para planificar rutas de recolección de residuos sobre una red vial real. "
        "La solución integra datos geográficos, nodos de clientes, puntos operativos, estimación de demanda, compactación espacial, "
        "optimización de recorridos, asignación de viajes a camiones y visualización interactiva en mapas HTML. "
        "Se trabajaron dos escenarios de operación: 60 toneladas y 120 toneladas de residuos."
    )
    add_paragraph(
        doc,
        "En el escenario de 60 toneladas se asignaron 6 viajes usando 4 camiones, con una distancia total de "
        f"{fmt_num(esc60['global']['distancia_total_km'])} km y un tiempo total de flota de "
        f"{fmt_num(esc60['global']['tiempo_total_h'])} h. En el escenario de 120 toneladas se asignaron 13 viajes usando "
        f"{esc120['global']['camiones_usados']} camiones, con {fmt_num(esc120['global']['distancia_total_km'])} km y "
        f"{fmt_num(esc120['global']['tiempo_total_h'])} h. Ambos escenarios usan la misma base de "
        f"{esc120['global']['clientes']} clientes operativos compactados y una capacidad por camión de "
        f"{fmt_num(esc120['global']['capacidad_camion_kg'], 0)} kg."
    )

    add_heading(doc, "2. Objetivo del proyecto")
    add_paragraph(
        doc,
        "El objetivo general fue construir una herramienta capaz de proponer rutas operativas de recolección que respeten "
        "restricciones de capacidad, jornada laboral, velocidades diferenciadas por etapa y ubicación de estación/relleno. "
        "Además, se buscó entregar resultados verificables mediante archivos geográficos, tablas de indicadores y mapas interactivos."
    )
    add_bullets(
        doc,
        [
            "Modelar la red vial y los clientes como un grafo ruteable.",
            "Distribuir y escalar la demanda de residuos para escenarios de 60 t y 120 t.",
            "Compactar clientes cercanos para representar paradas operativas realistas.",
            "Optimizar rutas mediante clustering, Clarke & Wright y mejora local 2-opt.",
            "Asignar viajes a una flota limitada respetando capacidad de 12,000 kg y jornada de 8 h.",
            "Generar salidas para análisis: HTML, GeoPackage, CSV y KPIs operativos.",
        ],
    )

    add_heading(doc, "3. Datos y herramientas utilizadas")
    add_paragraph(
        doc,
        "La implementación se desarrolló en Python y utiliza librerías de análisis geoespacial, redes y datos tabulares. "
        "Los archivos principales del proyecto permiten reproducir el proceso desde la preparación del grafo hasta la visualización final."
    )
    add_table(
        doc,
        ["Recurso", "Uso dentro del proyecto"],
        [
            ["grafo_cuenca.graphml", "Grafo base de la red vial."],
            ["grafo_actualizado.graphml", "Grafo con nodos de clientes, estación y relleno preparados para ruteo."],
            ["nodosAreaFinal.gpkg", "Capa geográfica de nodos seleccionados como clientes."],
            ["clientes_compactados.gpkg", "Salida del pipeline de compactación usada como base operativa."],
            ["ecu_pd_2020_1km.tif", "Raster de densidad poblacional usado para enriquecer nodos."],
            ["main.py", "Script principal de cálculo, optimización, exportación y visualización."],
            ["pipeline_compactación.py", "Proceso previo de compactación de clientes cercanos y callejones."],
            ["gpkg_to_graphml.py", "Actualiza el GraphML marcando clientes desde una capa GPKG."],
            ["Leaflet", "Visualización web interactiva de rutas, capas y animación."],
            ["OSMnx, NetworkX", "Carga del grafo y cálculo de rutas mínimas sobre la red vial."],
            ["GeoPandas, Shapely", "Manejo de capas geográficas y geometrías."],
            ["Pandas, NumPy", "Cálculo y exportación de indicadores."],
        ],
    )

    add_heading(doc, "4. Metodología implementada")
    add_paragraph(
        doc,
        "El flujo metodológico parte de la preparación del grafo vial y termina con una salida visual y tabular. "
        "La red se usa como base para calcular distancias reales por calles, no distancias euclidianas simples."
    )
    add_bullets(
        doc,
        [
            "Carga y validación del grafo vial: se verifica que las aristas tengan longitud y que los nodos de estación, relleno y clientes sean accesibles.",
            "Preparación de clientes: los nodos de atención se marcan en el grafo y se compactan para representar paradas operativas.",
            "Escalamiento de demanda: se ajusta la basura total objetivo a 60,000 kg o 120,000 kg según el escenario.",
            "Matriz origen-destino: se calculan distancias mínimas entre estación, relleno y clientes compactados.",
            "Zonificación: se evalúan distintos valores de k para dividir el territorio en zonas operativas.",
            "Optimización de rutas: se aplica Clarke & Wright dentro de zonas y luego una mejora 2-opt para reducir distancia interna.",
            "Restricciones operativas: se consideran capacidad máxima por viaje, jornada de 8 h, velocidades por etapa y tiempo por parada.",
            "Asignación a camiones: los viajes se distribuyen entre vehículos disponibles buscando cumplir tiempos y capacidades.",
            "Exportación: se generan mapas HTML, rutas por camión en GPKG y tablas CSV de indicadores.",
        ],
    )

    add_heading(doc, "5. Parámetros operativos")
    add_table(
        doc,
        ["Parámetro", "Valor usado"],
        [
            ["Capacidad máxima por camión", "12,000 kg"],
            ["Camiones máximos disponibles", "10"],
            ["Jornada máxima por camión", "8 h"],
            ["Recolectores por camión", "3"],
            ["Choferes por camión", "1"],
            ["Tiempo de parada", "30 s"],
            ["Velocidad estación -> primer cliente", "50 km/h"],
            ["Velocidad durante recolección", "10 km/h"],
            ["Velocidad último cliente -> relleno", "40 km/h"],
            ["Velocidad relleno -> estación", "50 km/h"],
            ["Precio diesel", "$2.99 por galón"],
            ["Rendimiento estimado", "4.5 km/galón"],
        ],
    )

    add_heading(doc, "6. Resultados globales por escenario")
    add_table(
        doc,
        ["Indicador", "60 toneladas", "120 toneladas"],
        [
            ["Clientes/paradas operativas", esc60["global"]["clientes"], esc120["global"]["clientes"]],
            ["Carga total asignada (kg)", fmt_num(esc60["global"]["carga_total_kg"], 0), fmt_num(esc120["global"]["carga_total_kg"], 0)],
            ["Camiones usados", esc60["global"]["camiones_usados"], esc120["global"]["camiones_usados"]],
            ["Camiones máximos disponibles", esc60["global"]["camiones_maximos"], esc120["global"]["camiones_maximos"]],
            ["Viajes asignados", esc60["global"]["viajes_asignados"], esc120["global"]["viajes_asignados"]],
            ["Distancia total (km)", fmt_num(esc60["global"]["distancia_total_km"]), fmt_num(esc120["global"]["distancia_total_km"])],
            ["Tiempo total flota (h)", fmt_num(esc60["global"]["tiempo_total_h"]), fmt_num(esc120["global"]["tiempo_total_h"])],
            ["Eficiencia (t/km)", fmt_num(esc60["global"]["eficiencia_ton_km"], 4), fmt_num(esc120["global"]["eficiencia_ton_km"], 4)],
            ["Costo diesel estimado", fmt_money(esc60["global"]["costo_diesel_usd"]), fmt_money(esc120["global"]["costo_diesel_usd"])],
            ["Costo laboral proporcional", fmt_money(esc60["global"]["costo_laboral_usd"]), fmt_money(esc120["global"]["costo_laboral_usd"])],
            ["Costo operativo estimado", fmt_money(esc60["global"]["costo_operativo_usd"]), fmt_money(esc120["global"]["costo_operativo_usd"])],
        ],
    )
    add_paragraph(
        doc,
        "La duplicación de la demanda de 60 t a 120 t no duplicó exactamente los camiones usados: pasó de 4 a 5 camiones. "
        "Sin embargo, sí aumentó de forma importante el número de viajes, de 6 a 13, porque la capacidad por viaje se mantiene "
        "limitada a 12,000 kg y algunas rutas finales quedan condicionadas por la geometría de la zona y la jornada."
    )

    add_heading(doc, "7. Resultados por camión")
    for esc in [esc60, esc120]:
        add_heading(doc, esc["nombre"], level=2)
        rows = []
        for c in esc["camiones"]:
            rows.append(
                [
                    c["id"],
                    c["viajes"],
                    fmt_num(c["carga_kg"], 2),
                    fmt_num(c["distancia_km"], 2),
                    fmt_num(c["tiempo_h"], 2),
                    fmt_num(c["uso_jornada_pct"], 2) + "%",
                    fmt_money(c["costo_operativo_usd"]),
                ]
            )
        add_table(
            doc,
            ["Camión", "Viajes", "Carga kg", "Dist. km", "Tiempo h", "Uso jornada", "Costo operativo"],
            rows,
        )

    add_heading(doc, "8. Resultados por viaje")
    for esc in [esc60, esc120]:
        add_heading(doc, esc["nombre"], level=2)
        rows = []
        for v in esc["viajes"]:
            util = 100.0 * float(v["carga_kg"]) / float(v["capacidad_kg"])
            rows.append(
                [
                    f"{v['camion']}.{v['viaje']}",
                    v["origen"],
                    fmt_num(v["carga_kg"], 2),
                    fmt_num(km_viaje(v), 2),
                    fmt_num(min_viaje(v), 1),
                    fmt_num(util, 2) + "%",
                ]
            )
        add_table(doc, ["Viaje", "Origen", "Carga kg", "Dist. km", "Tiempo min", "Utilización"], rows)

    add_heading(doc, "9. Comparación de algoritmos y selección de zonas")
    add_paragraph(
        doc,
        "Para el escenario de 120 t se comparó una base Clarke & Wright por distancia contra la versión híbrida con clustering, "
        "Clarke & Wright por zona y mejora 2-opt. La versión híbrida redujo el tiempo estimado frente a la base, aunque aumentó "
        "la distancia total en la comparación algorítmica previa a la asignación operativa."
    )
    add_table(
        doc,
        ["Algoritmo", "Rutas", "Paradas", "Carga kg", "Dist. km", "Tiempo h", "kg/km", "Mejora tiempo vs base"],
        [
            [
                row["algoritmo"],
                int(row["rutas"]),
                int(row["paradas"]),
                fmt_num(row["carga_kg"], 0),
                fmt_num(row["distancia_km"]),
                fmt_num(row["tiempo_h"]),
                fmt_num(row["kg_por_km"]),
                "-" if pd.isna(row["mejora_tiempo_pct_vs_base"]) else fmt_num(row["mejora_tiempo_pct_vs_base"]) + "%",
            ]
            for _, row in comparacion_alg.iterrows()
        ],
    )
    add_paragraph(
        doc,
        "La selección automática de zonas evaluó k entre 5 y 10. El menor score correspondió a k=5, por lo que se eligieron "
        "5 zonas para equilibrar distancia, tiempo, cantidad de viajes estimados y dispersión de cargas."
    )
    add_table(
        doc,
        ["k", "Rutas C&W", "Viajes estimados", "% baja carga", "Dist. km", "Tiempo h", "Score"],
        [
            [
                int(row["k"]),
                int(row["rutas_cw"]),
                int(row["viajes_estimados_capacidad"]),
                fmt_num(100 * row["pct_viajes_baja_carga"]) + "%",
                fmt_num(row["distancia_km"]),
                fmt_num(row["tiempo_h"]),
                fmt_num(row["score_k"], 3),
            ]
            for _, row in comparacion_k.iterrows()
        ],
    )

    add_heading(doc, "10. Visualización y archivos generados")
    add_paragraph(
        doc,
        "El proyecto genera una visualización HTML con Leaflet donde se pueden activar capas de clientes, puntos clave, calles base, "
        "rutas completas y viajes individuales. También incluye animación de camiones y paneles de indicadores."
    )
    add_table(
        doc,
        ["Archivo", "Contenido"],
        [
            ["mapa_rutas_recoleccion_60t.html", "Mapa interactivo del escenario de 60 t."],
            ["mapa_rutas_recoleccion_120t.html", "Mapa interactivo del escenario de 120 t."],
            ["rutas_animadas.html", "Visualización animada del escenario de 60 t."],
            ["rutas_animadas_120t.html", "Visualización animada del escenario de 120 t."],
            ["rutas_tramos_Camion_X.gpkg", "Tramos geográficos por camión para abrir en QGIS."],
            ["base_clientes.gpkg", "Clientes con demanda, densidad poblacional y distancias a puntos clave."],
            ["base_calles.gpkg", "Red vial base exportada."],
            ["base_puntos_clave.gpkg", "Estación y relleno sanitario."],
            ["kpi_camiones.csv", "Indicadores por camión de la última corrida exportada."],
            ["kpi_viajes.csv", "Indicadores por viaje de la última corrida exportada."],
            ["comparacion_algoritmos.csv", "Comparación entre enfoques de optimización."],
            ["comparacion_k_zonas.csv", "Evaluación de cantidad de zonas."],
        ],
    )

    add_heading(doc, "11. Interpretación de resultados")
    add_paragraph(
        doc,
        "El escenario de 60 t representa una operación más liviana: permite cubrir la demanda con 4 camiones y 6 viajes, pero un camión "
        "queda muy cercano al límite de jornada, con 99.91% de uso. Esto indica que, aunque la flota alcanza, existe poca holgura para "
        "imprevistos en esa ruta específica."
    )
    add_paragraph(
        doc,
        "El escenario de 120 t exige mayor coordinación: usa 5 camiones y 13 viajes. Los camiones 1 y 2 concentran la mayor carga "
        "y se acercan más al límite operativo, mientras que los camiones 4 y 5 trabajan con menor utilización de capacidad. Esto se explica "
        "porque algunas rutas largas acumulan muchas paradas de baja carga, lo cual consume tiempo de recolección aunque no complete el peso máximo."
    )
    add_paragraph(
        doc,
        "La eficiencia en toneladas por kilómetro aumenta de 0.2104 t/km en 60 t a 0.2330 t/km en 120 t. Esto sugiere que, al aumentar "
        "la demanda, una parte del recorrido fijo hacia relleno y estación se aprovecha mejor, aunque a costa de más viajes y más costo total."
    )

    add_heading(doc, "12. Limitaciones identificadas")
    add_bullets(
        doc,
        [
            "Las demandas dependen de reglas de estimación y escalamiento; sería ideal validarlas con datos históricos reales por sector.",
            "La velocidad se modela por etapa operativa y no con tráfico horario real.",
            "La asignación no considera todavía ventanas horarias de recolección por barrio, restricciones por ruido o restricciones municipales específicas.",
            "El costo operativo incluye diesel y costo laboral proporcional, pero no mantenimiento, depreciación, seguros ni peajes.",
            "Algunos archivos CSV corresponden a la última corrida exportada; para futuras pruebas se recomienda guardar salidas con sufijo de escenario.",
            "La relación entre demanda y densidad poblacional resultó baja en la corrida registrada, por lo que conviene revisar variables explicativas adicionales.",
        ],
    )

    add_heading(doc, "13. Trabajo futuro y continuidad del proyecto")
    add_paragraph(
        doc,
        "Para que otra persona pueda continuar el proyecto, se recomienda trabajar en mejoras incrementales y documentadas. "
        "La prioridad debe ser reproducibilidad, validación con datos reales y robustecimiento del modelo operativo."
    )
    add_table(
        doc,
        ["Línea futura", "Qué se puede implementar", "Beneficio esperado"],
        [
            ["Versionado de salidas", "Generar kpi_viajes_60t.csv, kpi_viajes_120t.csv y equivalentes por camión.", "Evita confundir resultados entre escenarios."],
            ["Validación de demanda", "Comparar demandas simuladas contra registros municipales o pesajes reales.", "Aumenta confiabilidad del modelo."],
            ["Tráfico por hora", "Incorporar velocidades diferenciadas por hora pico y tipo de vía.", "Rutas más realistas."],
            ["Costos completos", "Agregar mantenimiento, depreciación, llantas, seguros y costo por tonelada.", "Mejor evaluación económica."],
            ["Optimización multiobjetivo", "Balancear simultáneamente distancia, tiempo, costo, emisiones y equidad de carga.", "Mejor calidad operativa."],
            ["Interfaz de usuario", "Crear una app web o panel para seleccionar escenarios sin consola.", "Facilita uso por operadores no técnicos."],
            ["Reportes automáticos", "Exportar Word/PDF con métricas, mapas y tablas al terminar cada corrida.", "Ahorra tiempo de documentación."],
            ["Validación en campo", "Comparar rutas propuestas contra recorridos GPS reales.", "Detecta desvíos, calles no transitables y tiempos de servicio reales."],
            ["Restricciones viales", "Agregar sentidos, pendientes, calles angostas, zonas prohibidas y horarios.", "Reduce errores de implementación real."],
            ["Reoptimización dinámica", "Permitir recalcular ante camiones no disponibles, bloqueos o cambios de demanda.", "Aumenta resiliencia operativa."],
        ],
    )

    add_heading(doc, "14. Guía rápida para reproducir")
    add_paragraph(doc, "Desde una consola en la carpeta del proyecto se puede ejecutar:")
    doc.add_paragraph("python main.py --escenario 60t", style="Intense Quote")
    doc.add_paragraph("python main.py --escenario 120t", style="Intense Quote")
    add_paragraph(
        doc,
        "También se pueden ajustar parámetros como capacidad de camión, cantidad de camiones, jornada, velocidades, sueldos, diesel y rendimiento. "
        "El archivo GUIA_EJECUCION_Y_VISUALIZACION.md contiene la explicación completa de argumentos y lectura del mapa."
    )

    add_heading(doc, "15. Conclusiones")
    add_bullets(
        doc,
        [
            "El proyecto logró construir un flujo completo de optimización y visualización de rutas de recolección.",
            "Se generaron resultados para dos escenarios: 60 t y 120 t, ambos con 951 clientes/paradas operativas.",
            "El escenario de 60 t requiere 4 camiones y 6 viajes; el de 120 t requiere 5 camiones y 13 viajes.",
            "El aumento de demanda mejora la eficiencia ton/km, pero eleva distancia total, tiempo de flota y costo operativo.",
            "La solución ya es útil para análisis académico y prototipado operativo, pero debe fortalecerse con datos reales, versionado de salidas y restricciones de campo.",
        ],
    )

    add_heading(doc, "16. Anexo: lectura de los mapas")
    add_bullets(
        doc,
        [
            "Las líneas grises representan el recorrido total planificado.",
            "La línea oscura muestra la ruta dibujada durante la animación.",
            "Los marcadores R1, R2, etc. representan camiones animados.",
            "El código 2.3 significa camión 2, viaje 3.",
            "Cada viaje sigue el flujo: estación o relleno -> recolección -> relleno.",
            "El control de capas permite activar clientes, puntos clave, calles, rutas por camión y viajes individuales.",
        ],
    )

    doc.save(OUT_DOCX)
    print(f"Informe generado: {OUT_DOCX}")


if __name__ == "__main__":
    crear_informe()
