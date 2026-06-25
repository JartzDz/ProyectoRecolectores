from pathlib import Path
import json
import re

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parent
OUT_DOCX = BASE_DIR / "Informe_Final_Redes_Cuenca.docx"
LOGO = BASE_DIR / "docx_media_base_unzip" / "word" / "media" / "image1.jpg"
FLOWCHART = BASE_DIR / "docx_media_base_unzip" / "word" / "media" / "image3.png"
SUELDO_RECOLECTOR_USD = 665.0
SUELDO_CHOFER_USD = 801.0
RECOLECTORES_POR_CAMION = 3
CHOFERES_POR_CAMION = 1
DIAS_TRABAJO_MES = 30


def costo_laboral_diario_camion():
    return (
        RECOLECTORES_POR_CAMION * SUELDO_RECOLECTOR_USD
        + CHOFERES_POR_CAMION * SUELDO_CHOFER_USD
    ) / DIAS_TRABAJO_MES


def extract_const_from_html(path: Path, name: str):
    txt = path.read_text(encoding="utf-8", errors="ignore")
    match = re.search(rf"const\s+{name}\s*=\s*", txt)
    if not match:
        raise ValueError(f"No se encontro const {name} en {path.name}")

    start = match.end()
    depth = 0
    in_string = False
    escape = False

    for pos in range(start, len(txt)):
        ch = txt[pos]
        if in_string:
            if escape:
                escape = False
            elif ch == "\\":
                escape = True
            elif ch == '"':
                in_string = False
        else:
            if ch == '"':
                in_string = True
            elif ch in "[{":
                depth += 1
            elif ch in "]}":
                depth -= 1
            elif ch == ";" and depth == 0:
                return json.loads(txt[start:pos])
    raise ValueError(f"No se pudo cerrar const {name} en {path.name}")


def fmt_num(value, digits=2):
    try:
        return f"{float(value):,.{digits}f}"
    except (TypeError, ValueError):
        return str(value)


def fmt_money(value):
    return f"${fmt_num(value, 2)}"


def fmt_hm(hours):
    total_min = int(round(float(hours) * 60))
    return f"{total_min // 60}h {total_min % 60:02d}m"


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def shade_cell(cell, fill="D9EAF7"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(10)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    set_repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        set_cell_text(header_cells[idx], header, bold=True)
        shade_cell(header_cells[idx])
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
    doc.add_paragraph()
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {min(level, 3)}"
    run = p.add_run(text)
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.bold = True
    run.font.size = Pt(12)
    return p


def add_paragraph(doc, text, align=None):
    p = doc.add_paragraph()
    p.alignment = align or WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    r.font.name = "Arial"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    r.font.size = Pt(12)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(item)
        r.font.name = "Arial"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        r.font.size = Pt(12)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    for style_name in ["Normal", "List Bullet"]:
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(12)
    styles["Normal"].paragraph_format.line_spacing = 1.15
    styles["Normal"].paragraph_format.space_after = Pt(6)


def add_cover(doc):
    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO), width=Inches(4.8))

    for line, size, bold in [
        ("FACULTAD DE INGENIERIA", 14, True),
        ("Carrera de Computacion", 12, False),
        ("Redes de Computadoras", 12, False),
        ("Septimo Semestre", 12, False),
        ("Optimización de rutas de recolección de desechos sólidos en la ciudad de Cuenca", 16, True),
        ("Algoritmo heurístico Clarke & Wright con zonificación y mejora 2-opt", 13, True),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        r.bold = bold
        r.font.name = "Arial"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        r.font.size = Pt(size)

    doc.add_paragraph()
    add_paragraph(doc, "Estudiantes: completar con los nombres del grupo", WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, "Docente: Ing. Raul Ortiz Gaona, PhD", WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, "Cuenca, 13 de junio de 2026", WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()


def global_rows(metrics60, metrics120):
    rows = []
    for label, metrics in [("60", metrics60), ("120", metrics120)]:
        g = metrics["global"]
        costo_trabajadores = int(g["camiones_usados"]) * costo_laboral_diario_camion()
        costo_total = costo_trabajadores + float(g["costo_diesel_usd"])
        rows.append([
            int(g["viajes_asignados"]),
            int(g["camiones_usados"]),
            label,
            fmt_num(g["distancia_total_km"], 2),
            fmt_num(g["eficiencia_ton_km"], 4),
            fmt_hm(g["tiempo_total_h"]),
            int(g["camiones_usados"]) * 4,
            fmt_money(costo_trabajadores),
            fmt_money(g["costo_diesel_usd"]),
            fmt_money(costo_total),
        ])
    return rows


def build_report():
    doc = Document()
    configure_document(doc)

    metrics60 = extract_const_from_html(BASE_DIR / "mapa_rutas_recoleccion_60t.html", "METRICAS")
    metrics120 = extract_const_from_html(BASE_DIR / "mapa_rutas_recoleccion_120t.html", "METRICAS")
    comparacion_alg = pd.read_csv(BASE_DIR / "comparacion_algoritmos.csv")
    comparacion_k = pd.read_csv(BASE_DIR / "comparacion_k_zonas.csv")
    kpi_camiones = pd.read_csv(BASE_DIR / "kpi_camiones.csv")
    kpi_rutas = pd.read_csv(BASE_DIR / "kpi_viajes.csv")

    add_cover(doc)

    add_heading(doc, "1. Antecedentes")
    add_paragraph(doc, "La recolección de desechos sólidos urbanos es un servicio esencial para la salud pública, la limpieza de calles y la protección ambiental. En una ciudad, el servicio debe coordinar puntos de generación de residuos, estaciones de salida, relleno sanitario, personal, camiones, horarios de atención y una red vial con sentidos, distancias y velocidades variables. Cuando las rutas se definen sin soporte técnico, aparecen recorridos repetidos, rutas con baja utilización, mayor consumo de combustible y tiempos improductivos.")
    add_paragraph(doc, "El problema puede formularse como un Vehicle Routing Problem (VRP), donde un conjunto de vehículos debe visitar puntos de demanda y regresar o descargar en un destino, minimizando distancia, tiempo o costo y respetando restricciones operativas. Debido a que el VRP crece de forma combinatoria, en escenarios urbanos reales se suelen usar algoritmos heurísticos o metaheurísticos que entregan soluciones suficientemente buenas en tiempos de cálculo aceptables.")

    add_heading(doc, "2. Situación en la ciudad de Cuenca")
    add_paragraph(doc, "Cuenca presenta una red vial urbana con zonas consolidadas, vías de diferente jerarquía, sectores con calles estrechas y puntos de alta concentración de actividades. Estas condiciones dificultan la planificación manual de rutas de recolección, porque no basta con unir puntos por cercanía geográfica: se requiere considerar la distancia real por calles, la conectividad del grafo vial, la ubicación de estaciones y relleno sanitario, y la distribución espacial de la demanda.")
    add_paragraph(doc, "En el proyecto se modeló la ciudad mediante información georreferenciada y un grafo vial dirigido. Los nodos representan puntos de la red o paradas de recolección, mientras que las aristas representan tramos de calle con longitud. Esta representación permite calcular recorridos sobre calles reales y evaluar indicadores operativos como distancia, tiempo, carga recogida, eficiencia, uso de flota y costos estimados.")
    add_paragraph(doc, "Durante la validación se detectaron problemas propios de trabajar con redes viales abiertas: componentes desconectadas, clientes inaccesibles desde la estación, clientes sin ruta hacia el relleno, calles sin nombre y aristas sin geometría completa. Estos problemas fueron tratados mediante auditoría del grafo, exclusión de nodos no conectados, corrección de distancias y compactación de nodos operativos.")

    add_heading(doc, "3. Justificación del proyecto")
    add_paragraph(doc, "La realización del proyecto beneficia a la ciudad porque permite comparar rutas de recolección con criterios cuantificables y reproducibles. Un sistema informático de este tipo ayuda a reducir recorridos innecesarios, balancear la carga entre camiones, estimar costos de operación y visualizar rutas en mapas interactivos. Además, facilita que futuras decisiones municipales se apoyen en datos y no únicamente en experiencia operativa.")
    add_paragraph(doc, "El análisis se mantuvo siempre en dos escenarios: 60 toneladas y 120 toneladas. Estos valores permiten observar el comportamiento del sistema cuando la demanda diaria se mantiene en un volumen moderado y cuando se duplica, de modo que la comparación no dependa de una sola corrida.")
    add_paragraph(doc, "Desde el punto de vista académico, el proyecto integra redes de computadoras, análisis de grafos, datos geoespaciales, algoritmos heurísticos y programación en Python. El resultado no es solo una ruta dibujada, sino un flujo completo que transforma datos de ciudad en métricas comparables para tomar decisiones.")

    add_heading(doc, "4. Objetivo general")
    add_paragraph(doc, "Desarrollar un sistema informático, implementando un algoritmo heurístico, que determine las rutas que deben recorrer los vehículos recolectores de residuos sólidos en la Ciudad de Cuenca.")

    add_heading(doc, "5. Objetivos específicos")
    add_bullets(doc, [
        "Identificar los algoritmos de enrutamiento para Vehicle Routing Problem (VRP), utilizando diferentes plataformas de inteligencia artificial.",
        "Definir criterios que permitan evaluar y seleccionar los algoritmos más eficientes para el caso de recolección de residuos.",
        "Desarrollar aplicaciones informáticas en lenguaje Python, implementando los algoritmos seleccionados.",
        "Ejecutar las aplicaciones y determinar cuáles algoritmos arrojan los mejores resultados según métricas operativas.",
    ])

    add_heading(doc, "6. Metodología")
    add_heading(doc, "6.1. Definición de términos utilizados", 2)
    add_table(doc, ["Término", "Definición aplicada en el proyecto"], [
        ["VRP", "Problema de ruteo de vehículos: busca planificar recorridos para atender puntos de demanda con una flota limitada."],
        ["Grafo vial", "Representación de la ciudad con nodos e intersecciones unidos por aristas o tramos de calle."],
        ["Nodo cliente", "Punto o parada operativa donde se recoge una cantidad estimada de residuos."],
        ["Ruta", "Recorrido operativo completo: el camión inicia el recorrido, recolecta residuos hasta llenarse o completar su asignación, y descarga en el depósito o relleno sanitario."],
        ["Demanda", "Volumen o peso de residuos asignado a un nodo o sector de la ciudad."],
        ["Métrica", "Indicador usado para evaluar el resultado: distancia, tiempo, costo, carga, eficiencia, rutas y flota."],
    ])

    add_heading(doc, "6.2. Algoritmo heurístico", 2)
    add_paragraph(doc, "Un algoritmo heurístico es un método que busca soluciones buenas sin revisar todas las combinaciones posibles. En problemas como el VRP, una búsqueda exacta puede ser inviable cuando aumenta el número de clientes, por lo que las heurísticas construyen y mejoran soluciones usando reglas prácticas. La calidad del resultado se evalúa comparando métricas: menor distancia, menor tiempo, mejor uso de capacidad, menor costo y cumplimiento de restricciones.")
    add_paragraph(doc, "La heurística de Clarke & Wright parte de rutas simples y calcula ahorros al unir pares de clientes. Si una fusión reduce distancia y respeta las condiciones del problema, se acepta; si no, se descarta y se continúa con el siguiente ahorro. En este proyecto se usó una variante extendida con zonificación inicial y mejora local 2-opt.")

    add_heading(doc, "6.3. Estrategia de búsqueda usando plataformas IA", 2)
    add_paragraph(doc, "Se consultaron plataformas de inteligencia artificial como ChatGPT, DeepSeek, Claude, Gemini y Copilot. La estrategia consistió en usar dos tipos de prompts: uno general, para identificar algoritmos VRP ampliamente recomendados, y otro específico, orientado a rutas de recolección de basura con restricciones de capacidad, tiempo, puntos de salida y descarga, costos y operación urbana. Luego se consolidaron las respuestas repetidas y se contrastaron con criterios técnicos.")

    add_heading(doc, "6.4. Preselección de diez algoritmos recomendados", 2)
    add_table(doc, ["Algoritmo", "Descripción resumida"], [
        ["Clarke & Wright Savings", "Heurística clásica que fusiona rutas según el ahorro de distancia."],
        ["Algoritmo Genético", "Metaheurística evolutiva que combina y muta soluciones candidatas."],
        ["Búsqueda Tabú", "Explora vecindarios evitando regresar a soluciones recientes mediante una lista tabú."],
        ["Ant Colony Optimization", "Construye rutas usando feromonas artificiales inspiradas en colonias de hormigas."],
        ["ALNS", "Destruye y reconstruye partes de la solución con operadores adaptativos."],
        ["Simulated Annealing", "Acepta ocasionalmente soluciones peores para escapar de óptimos locales."],
        ["Nearest Neighbor", "Construye rutas eligiendo iterativamente el cliente más cercano."],
        ["Sweep Algorithm", "Ordena clientes angularmente y forma rutas por barrido geográfico."],
        ["Branch and Bound / MILP", "Enfoque exacto que busca optimalidad con modelos matemáticos y poda de ramas."],
        ["Memetic Algorithm", "Combina algoritmos genéticos con búsqueda local para refinar individuos."],
    ])

    add_heading(doc, "6.5. Criterios para elegir los tres mejores algoritmos", 2)
    add_table(doc, ["Criterio", "Peso (%)", "Sentido de evaluación"], [
        ["Calidad de solución", "33.00", "Capacidad de reducir distancia, tiempo y costo."],
        ["Robustez", "15.25", "Resistencia ante restricciones urbanas y operativas."],
        ["Determinismo", "13.00", "Repetibilidad de resultados bajo las mismas condiciones."],
        ["Sensibilidad de hiperparámetros", "11.75", "Estabilidad frente a ajustes de parámetros."],
        ["Escalabilidad", "11.00", "Capacidad de trabajar con grafos y demandas grandes."],
        ["Facilidad de implementación", "10.75", "Complejidad de programación, librerías y hardware."],
        ["Interpretabilidad", "5.25", "Facilidad para explicar cómo se obtuvo la ruta."],
    ])

    add_heading(doc, "6.6. Identificación de los diez mejores algoritmos VRP", 2)
    add_paragraph(doc, "Luego de consolidar las recomendaciones y aplicar los criterios de evaluación, los algoritmos con mayor prioridad fueron Algoritmo Genético, Clarke & Wright y Búsqueda Tabú. Para el grupo se asignó Clarke & Wright, debido a su equilibrio entre interpretabilidad, facilidad de implementación y buen desempeño en escenarios de distribución y recolección.")

    add_heading(doc, "6.7. Restricciones de operación del servicio", 2)
    add_paragraph(doc, "Para esta sección se consideran restricciones propias del servicio y de la ciudad: los residuos deben ser atendidos sin omitir zonas, los recorridos deben seguir calles existentes, la conectividad del grafo debe permitir llegar a cada punto, las descargas deben terminar en el relleno sanitario y la planificación debe evitar rutas discontinuas o imposibles de ejecutar en campo. Los parámetros específicos de flota, capacidad y jornada se usan más adelante dentro de la implementación y las métricas.")

    add_heading(doc, "6.8. Descripción del algoritmo mediante diagrama de flujo", 2)
    add_paragraph(doc, "El flujo implementado inicia con la división territorial y asignación operativa. Después se selecciona una zona, se calculan caminos mínimos con Dijkstra, se crean rutas iniciales, se calculan ahorros entre pares de nodos y se ordenan de mayor a menor. Cada posible fusión se evalúa verificando extremos y capacidad; si es válida, se fusiona, y si no, se descarta. Al final se materializan las rutas optimizadas sobre el grafo vial.")
    if FLOWCHART.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(FLOWCHART), width=Inches(5.9))
        add_paragraph(doc, "Figura 1. Diagrama de flujo del algoritmo Clarke & Wright extendido para recolección de residuos.", WD_ALIGN_PARAGRAPH.CENTER)

    add_heading(doc, "6.9. Implementación en Python", 2)
    add_paragraph(doc, "La implementación se realizó en Python en el archivo main.py. El sistema carga la red vial, valida longitudes de aristas, prepara clientes, escala la demanda, calcula matrices origen-destino, ejecuta zonificación, aplica Clarke & Wright por zona, mejora rutas con 2-opt, divide o fusiona rutas según condiciones operativas y asigna rutas a camiones.")
    add_table(doc, ["Parámetro operativo", "Valor usado"], [
        ["Capacidad máxima por ruta", "12,000 kg"],
        ["Camiones disponibles", "10"],
        ["Jornada máxima por camión", "8 horas"],
        ["Personal por camión", "3 recolectores y 1 chofer"],
        ["Tiempo de parada", "30 segundos"],
        ["Rendimiento estimado", "4.5 km/gal"],
        ["Precio de diesel", "$2.99 por galón"],
    ])
    add_paragraph(doc, f"Para el informe económico se corrigió el cálculo de trabajadores a una lógica mensual usada en Ecuador. La cuadrilla por camión se compone de {RECOLECTORES_POR_CAMION} recolectores y {CHOFERES_POR_CAMION} chofer, con sueldos mensuales de {fmt_money(SUELDO_RECOLECTOR_USD)} y {fmt_money(SUELDO_CHOFER_USD)}, respectivamente. Por tanto, el costo diario de personal por camión se calcula como ({RECOLECTORES_POR_CAMION} x {fmt_money(SUELDO_RECOLECTOR_USD)} + {CHOFERES_POR_CAMION} x {fmt_money(SUELDO_CHOFER_USD)}) / {DIAS_TRABAJO_MES} días = {fmt_money(costo_laboral_diario_camion())}.")

    add_heading(doc, "6.10. IA utilizada para generar el programa Python", 2)
    add_paragraph(doc, "La inteligencia artificial se utilizó como apoyo técnico y documental, no como reemplazo de la validación del grupo. Las respuestas fueron revisadas, adaptadas al caso de Cuenca y contrastadas con los archivos georreferenciados, resultados CSV y mapas generados.")
    add_table(doc, ["Plataforma o modelo", "Lenguajes/formatos apoyados", "Uso dentro del proyecto"], [
        ["ChatGPT / Codex", "Python, Markdown, Word, CSV", "Apoyo en estructuración del informe, revisión de código, generación de tablas y explicación de resultados."],
        ["DeepSeek", "Python y pseudocódigo", "Consulta comparativa de algoritmos VRP y criterios de selección."],
        ["Claude", "Texto técnico y pseudocódigo", "Redacción y contraste de metodología, limitaciones y conclusiones."],
        ["Google Gemini", "Texto técnico", "Búsqueda de alternativas de algoritmos y validación conceptual."],
        ["Microsoft Copilot", "Python y documentación", "Apoyo puntual para depuración, explicación de funciones y mejoras de implementación."],
        ["Python", "pandas, datos geoespaciales, HTML/JavaScript", "Lenguaje principal de implementación, cálculo de rutas, KPIs y visualización interactiva."],
    ])

    add_heading(doc, "6.11. Adquisición de información georreferenciada", 2)
    add_paragraph(doc, "La información georreferenciada se obtuvo mediante archivos GeoPackage, GraphML, HTML y datos auxiliares del proyecto. Entre los archivos usados están grafo_cuenca.graphml, grafo_actualizado.graphml, nodosAreaFinal.gpkg, clientes_compactados.gpkg, base_calles.gpkg, base_clientes.gpkg y base_puntos_clave.gpkg. Además, el proyecto conserva una caché de datos de OpenStreetMap generada con Overpass.")

    add_heading(doc, "6.12. Pasos para obtener el grafo", 2)
    add_bullets(doc, [
        "Descargar o cargar la red vial de Cuenca en formato de grafo.",
        "Validar que cada arista tenga longitud para poder calcular distancias reales.",
        "Agregar o asociar nodos de clientes, estación y relleno sanitario.",
        "Compactar puntos cercanos para representar paradas operativas manejables.",
        "Calcular caminos mínimos entre puntos relevantes usando Dijkstra.",
        "Exportar el grafo actualizado y capas geográficas para visualización y análisis.",
    ])

    add_heading(doc, "6.13. Volumen de desechos generado", 2)
    add_paragraph(doc, "La estimación de generación de residuos se apoyó en estudios locales de la Universidad del Azuay sobre generación per cápita y composición de residuos sólidos en Cuenca. En particular, se usaron como referencia los trabajos de Fiallo Flor (2020) y Arévalo Vélez y Muñoz Pauta (2010), que documentan tasas de generación para fuentes como hogares, restaurantes, hoteles, instituciones educativas y mercados.")
    add_paragraph(doc, "Además de la revisión bibliográfica, se consideró información obtenida mediante entrevista operativa del proyecto, de la cual se definieron los dos escenarios de análisis: 60 toneladas y 120 toneladas. Estos escenarios no representan predicciones absolutas de toda la ciudad, sino volúmenes de prueba para comparar el comportamiento de la solución cuando la demanda se mantiene en un nivel base y cuando se duplica.")

    add_heading(doc, "6.14. Reparto del volumen en la ciudad", 2)
    add_paragraph(doc, "El volumen generado se distribuyó sobre los clientes o paradas compactadas de la ciudad. La asignación no fue uniforme: se tomó como base la generación de un hogar urbano y se aplicaron relaciones relativas para establecimientos con mayor generación, como bares, restaurantes, hoteles, instituciones educativas y mercados, de acuerdo con las fuentes locales de Cuenca (Fiallo Flor, 2020; Arévalo Vélez & Muñoz Pauta, 2010).")
    add_table(doc, ["Tipo de fuente", "Generación estimada usada como referencia", "Relación aproximada"], [
        ["Hogar urbano", "2.01 kg/día", "1.00 vez"],
        ["Bar grande", "8.16 kg/día", "4 veces"],
        ["Restaurante grande", "17.08 kg/día", "8.5 veces"],
        ["Hotel grande", "11.44 kg/día", "5.7 veces"],
        ["Escuela fiscal", "74.55 kg/día", "37 veces"],
        ["Universidad", "223.50 kg/día", "111 veces"],
        ["Mercado", "14,490.00 kg/día", "7,208 veces"],
    ])
    add_paragraph(doc, "Luego, la demanda se escaló hasta cerrar exactamente los escenarios de 60,000 kg y 120,000 kg. Esto permite comparar las mismas reglas de distribución espacial bajo dos cargas totales distintas.")

    add_heading(doc, "6.15. Métricas de eficiencia del servicio", 2)
    add_paragraph(doc, "Para aproximar el tiempo de manipulación de residuos se tomó como referencia el estudio de Moore et al. (2021), donde se reporta un tiempo promedio de 2.44 segundos para levantar una bolsa de 20 kg. A partir de esa relación se usó una aproximación lineal de 0.122 segundos por kilogramo recolectado, útil para estimar tiempos de atención cuando la demanda cambia por nodo.")
    add_table(doc, ["Métrica", "Uso"], [
        ["Número de rutas", "Mide cuántos recorridos operativos completos se requieren."],
        ["Distancia total y por camión", "Evalúa consumo de red vial y costo de transporte."],
        ["Tiempo total y por ruta", "Permite verificar jornadas y holguras."],
        ["Carga recogida", "Verifica cobertura del volumen objetivo."],
        ["Eficiencia ton/km", "Relaciona toneladas recolectadas con distancia recorrida."],
        ["Costo diesel", "Aproxima el costo energético del recorrido."],
        ["Costo laboral diario", "Estima el costo de la cuadrilla por camión según sueldo mensual dividido para días de trabajo."],
        ["Número de trabajadores", "Relaciona flota usada con personal requerido."],
    ])

    add_heading(doc, "6.16. Ejecución del programa", 2)
    add_paragraph(doc, "El programa se ejecuta desde consola con escenarios predefinidos. Para 60 toneladas se usa python main.py --escenario 60t y para 120 toneladas se usa python main.py --escenario 120t. La ejecución genera mapas HTML, rutas por camión en GeoPackage y archivos CSV con indicadores por ruta, por camión y comparaciones de algoritmos.")

    add_heading(doc, "6.17. Resultados obtenidos en cuadro de resumen", 2)
    add_table(doc, ["Rutas", "Camiones", "Carga (Ton)", "Distancia (km)", "Eficiencia (Ton/km)", "Jornada", "# trabajadores", "Costo trabajadores ($)", "Costo diesel ($)", "Costo total ($)"], global_rows(metrics60, metrics120))

    add_heading(doc, "6.18. Formato de presentación de resultados", 2)
    add_paragraph(doc, "Los resultados se presentan en tablas comparativas por escenario, por camión y por algoritmo. El formato resume primero los indicadores globales y luego detalla la distribución operativa de camiones, rutas, carga, distancia, tiempo, costos y utilización.")

    add_heading(doc, "Resultados por camión del escenario 120 toneladas", 3)
    add_table(doc, ["Camión", "Rutas", "Carga (kg)", "Distancia (km)", "Tiempo", "Uso jornada", "Costo total diario"], [
        [
            int(row["camion"]),
            int(row["viajes_asignados"]),
            fmt_num(row["carga_total_kg"], 2),
            fmt_num(row["dist_total_km"], 2),
            fmt_hm(row["tiempo_ruta_h"]),
            fmt_num(row["uso_jornada_pct"], 2) + "%",
            fmt_money(float(row["costo_diesel_usd"]) + costo_laboral_diario_camion()),
        ]
        for _, row in kpi_camiones.iterrows()
    ])

    costo_trab_60 = metrics60["global"]["camiones_usados"] * costo_laboral_diario_camion()
    costo_trab_120 = metrics120["global"]["camiones_usados"] * costo_laboral_diario_camion()
    costo_total_60 = costo_trab_60 + metrics60["global"]["costo_diesel_usd"]
    costo_total_120 = costo_trab_120 + metrics120["global"]["costo_diesel_usd"]

    add_heading(doc, "6.19. Análisis de resultados", 2)
    add_paragraph(doc, f"En el escenario de 60 toneladas se requieren {metrics60['global']['viajes_asignados']} rutas y {metrics60['global']['camiones_usados']} camiones, con una distancia total de {fmt_num(metrics60['global']['distancia_total_km'])} km y eficiencia de {fmt_num(metrics60['global']['eficiencia_ton_km'], 4)} ton/km. En el escenario de 120 toneladas se requieren {metrics120['global']['viajes_asignados']} rutas y {metrics120['global']['camiones_usados']} camiones, con {fmt_num(metrics120['global']['distancia_total_km'])} km y eficiencia de {fmt_num(metrics120['global']['eficiencia_ton_km'], 4)} ton/km.")
    add_paragraph(doc, "Al duplicar la carga de 60 a 120 toneladas, la distancia y el costo aumentan, pero la eficiencia ton/km también mejora. Esto indica que parte del recorrido fijo se aprovecha mejor cuando los camiones transportan mayor volumen. Sin embargo, el escenario de 120 toneladas exige más coordinación, porque incrementa el número de rutas y la carga de asignación sobre la flota.")
    add_paragraph(doc, f"Con el criterio de sueldo mensual dividido para días de trabajo, el costo diario de trabajadores es {fmt_money(costo_trab_60)} para 60 toneladas y {fmt_money(costo_trab_120)} para 120 toneladas. Sumando diesel, el costo total estimado es {fmt_money(costo_total_60)} y {fmt_money(costo_total_120)}, respectivamente. Este cálculo es más realista para el contexto ecuatoriano que una tarifa estrictamente horaria.")

    add_heading(doc, "6.20. Comparación de los algoritmos implementados", 2)
    add_table(doc, ["Algoritmo", "Rutas", "Paradas", "Carga (kg)", "Distancia (km)", "Tiempo (h)", "kg/km", "Mejora tiempo"], [
        [
            row["algoritmo"],
            int(row["rutas"]),
            int(row["paradas"]),
            fmt_num(row["carga_kg"], 0),
            fmt_num(row["distancia_km"], 2),
            fmt_num(row["tiempo_h"], 2),
            fmt_num(row["kg_por_km"], 2),
            "-" if pd.isna(row["mejora_tiempo_pct_vs_base"]) else fmt_num(row["mejora_tiempo_pct_vs_base"], 2) + "%",
        ]
        for _, row in comparacion_alg.iterrows()
    ])
    add_paragraph(doc, "La comparación muestra que la variante híbrida con clustering inicial, Clarke & Wright por zona y 2-opt reduce el tiempo estimado frente a la base de Clarke & Wright por distancia. A cambio, recorre una distancia mayor, lo que refleja una decisión de balance operativo: se prioriza reducir tiempo y hacer más manejable la asignación por zonas.")
    add_table(doc, ["k", "Rutas C&W", "Rutas estimadas", "% baja carga", "Distancia (km)", "Tiempo (h)", "Score"], [
        [
            int(row["k"]),
            int(row["rutas_cw"]),
            int(row["viajes_estimados_capacidad"]),
            fmt_num(100 * row["pct_viajes_baja_carga"], 2) + "%",
            fmt_num(row["distancia_km"], 2),
            fmt_num(row["tiempo_h"], 2),
            fmt_num(row["score_k"], 3),
        ]
        for _, row in comparacion_k.iterrows()
    ])

    add_heading(doc, "6.21. Conclusiones, comentarios y recomendaciones", 2)
    add_heading(doc, "Problemas encontrados y solución aplicada", 3)
    add_table(doc, ["Problema", "Solución aplicada", "Resultado"], [
        ["Clientes inaccesibles o sin ruta al relleno", "Se validó conectividad del grafo y se excluyeron nodos sin conexión operativa.", "Las rutas finales se calcularon solo sobre nodos alcanzables."],
        ["Distancias y tiempos anómalos", "Se corrigió la red proyectando coordenadas y recalculando longitudes en metros.", "Los recorridos pasaron a tener distancias y tiempos más coherentes."],
        ["Demasiados puntos muy cercanos", "Se compactaron nodos en callejones y puntos del mismo lado de calle a menos de 80 m.", "Los clientes operativos bajaron a 944/951 según corrida, conservando la demanda total."],
        ["Costo laboral calculado por horas", "En el informe se corrigió a costo diario por cuadrilla, usando sueldos mensuales.", "El costo económico queda alineado con la forma de pago local."],
        ["Diferencias entre distancia y tiempo", "Se comparó Clarke & Wright base contra el híbrido con zonas y 2-opt.", "La variante híbrida mejora tiempo, aunque aumenta distancia."],
    ])

    add_heading(doc, "Logros alcanzados", 3)
    add_bullets(doc, [
        "Se implementó un flujo completo desde datos georreferenciados hasta rutas, KPIs y mapas interactivos.",
        "Se evaluaron de forma consistente los escenarios de 60 y 120 toneladas.",
        "Se integraron criterios reales de capacidad, jornada, diesel, personal y descarga en relleno.",
        "Se generaron tablas comparativas por escenario, por camión y por algoritmo.",
        "Se documentó el proceso con fuentes locales de Cuenca y referencias técnicas.",
    ])

    add_heading(doc, "Aspectos que no se lograron o quedaron limitados", 3)
    add_bullets(doc, [
        "No se validaron las rutas con GPS real de camiones recolectores en operación.",
        "No se incorporó tráfico por hora, cierres temporales, pendientes ni restricciones por ancho de vía.",
        "La demanda de residuos sigue siendo una estimación escalada, no una medición real por cuadra.",
        "No se modeló fatiga de trabajadores ni ergonomía de carga más allá del tiempo promedio por kilogramo.",
        "No se desarrolló una interfaz final para operadores municipales; la ejecución sigue dependiendo de consola y archivos.",
    ])

    add_heading(doc, "Trabajo a futuro", 3)
    add_table(doc, ["Línea futura", "Qué se debería hacer", "Referencia o justificación"], [
        ["Validación ergonómica", "Incluir distancia de contacto bolsa-cuerpo, postura y carga acumulada de trabajadores.", "Moore et al. (2021)."],
        ["Demanda real por sector", "Contrastar pesos estimados con registros de EMAC EP o pesajes reales por ruta.", "Fiallo Flor (2020); Arévalo Vélez & Muñoz Pauta (2010)."],
        ["Grafo vial robusto", "Cruzar OSM con capas municipales, Waze, OSRM, GraphHopper o servicios externos.", "Auditoría técnica de red vial del proyecto."],
        ["Tráfico y horarios", "Agregar velocidades por hora pico y ventanas de recolección por barrio.", "Mejora de realismo operativo."],
        ["Interfaz de usuario", "Crear una aplicación web para elegir escenario, revisar rutas y exportar reportes.", "Facilitar uso por personal no técnico."],
        ["Optimización multiobjetivo", "Balancear distancia, tiempo, costo, emisiones, carga y equidad laboral.", "Mejor toma de decisiones."],
    ])

    add_bullets(doc, [
        "Se logró construir un sistema informático en Python para generar rutas de recolección sobre una red vial georreferenciada de Cuenca.",
        "El algoritmo Clarke & Wright extendido permitió obtener rutas válidas y medibles para escenarios de 60 y 120 toneladas.",
        "La variante híbrida redujo el tiempo frente a la base de Clarke & Wright, aunque aumentó la distancia total estimada en la comparación algorítmica.",
        "El escenario de 120 toneladas requiere 13 rutas, 5 camiones y 20 trabajadores, con un costo diario estimado de " + fmt_money(costo_total_120) + ".",
        "Se recomienda validar la demanda con datos municipales reales, incorporar tráfico por horario y guardar salidas CSV separadas por escenario para evitar confusiones.",
        "Queda pendiente integrar restricciones de campo más específicas, como calles no transitables para camiones, pendientes, horarios barriales y ventanas de recolección.",
    ])

    add_heading(doc, "Referencias", 2)
    add_bullets(doc, [
        "Moore, S. M., et al. (2021). Exploring the prospective efficacy of waste bag-body contact allowance to reduce biomechanical exposure in municipal waste collection. Applied Ergonomics.",
        "Fiallo Flor, R. C. (2020). Análisis de la generación per cápita y composición gravimétrica de residuos sólidos procedentes de diferentes fuentes de la ciudad de Cuenca (Trabajo de titulación). Escuela de Ingeniería Civil y Gerencia de Construcciones, Facultad de Ciencia y Tecnología, Universidad del Azuay, Cuenca, Ecuador.",
        "Arévalo Vélez, C., & Muñoz Pauta, F. (2010). Evolución de las características de los residuos sólidos en el cantón Cuenca. Universidad del Azuay / Empresa Municipal de Aseo de Cuenca (EMAC EP), Cuenca, Ecuador.",
        "ETAPA EP. (2025). Publico/Vias MapServer. https://geo.etapa.net.ec/arcgis/rest/services/Publico/Vias/MapServer",
        "OpenStreetMap Wiki. (2025). Turn Restrictions; Forward & Backward, Left & Right. https://wiki.openstreetmap.org/",
    ])

    add_heading(doc, "Anexos")
    add_paragraph(doc, "Archivos principales generados: mapas HTML de los escenarios, CSV de indicadores por camión, CSV de indicadores por ruta, comparacion_algoritmos.csv y comparacion_k_zonas.csv.")
    add_paragraph(doc, f"El CSV de indicadores por ruta contiene {len(kpi_rutas)} rutas de la última corrida exportada.")

    try:
        doc.save(OUT_DOCX)
        return OUT_DOCX
    except PermissionError:
        fallback = BASE_DIR / "Informe_Final_Redes_Cuenca_actualizado.docx"
        doc.save(fallback)
        return fallback


if __name__ == "__main__":
    print(build_report())
